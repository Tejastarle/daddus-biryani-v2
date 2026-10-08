"""
Builds the customer downloads from data/menu.json:

  public/downloads/daddus-biryani-menu.pdf   — a designed menu card
  public/downloads/daddus-biryani-menu.xlsx  — the same data as a sheet

Run after scripts/extract_menu.py whenever prices change.
Fonts live in scripts/fonts/.
"""
import json, os
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (BaseDocTemplate, Frame, PageTemplate, Paragraph, Table,
                                TableStyle, Spacer, KeepTogether, NextPageTemplate,
                                CondPageBreak, Image as RLImage, PageBreak)
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_RIGHT, TA_CENTER
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FONTS = os.path.join(ROOT, 'scripts', 'fonts')
OUT = os.path.join(ROOT, 'public', 'downloads')
os.makedirs(OUT, exist_ok=True)
items = json.load(open(os.path.join(ROOT, 'data', 'menu.json')))

# Brand palette
GREEN = HexColor('#0E3B2F')
DEEP = HexColor('#082A21')
BRASS = HexColor('#C9A24B')
BRASS_DK = HexColor('#9A7A2E')
KESAR = HexColor('#E8912D')
IVORY = HexColor('#FBF5E9')
IVORY_DIM = HexColor('#F1E7D2')
INK = HexColor('#1F2A24')
MUTED = HexColor('#6B7A71')
FAINT = HexColor('#AEB8B1')
RULE = HexColor('#E0D4BA')
VEG, NONVEG = HexColor('#2E8B3E'), HexColor('#B3261E')

for name, f in [('Rozha', 'RozhaOne-Regular.ttf'), ('Mukta', 'Mukta-Regular.ttf'),
                ('Mukta-SB', 'Mukta-SemiBold.ttf'), ('Mukta-B', 'Mukta-Bold.ttf')]:
    pdfmetrics.registerFont(TTFont(name, os.path.join(FONTS, f)))

SECTIONS = [
    ('kolkata', 'Kolkata', 'Delicate and distinctive', 'Aloo and egg tradition', 'Mild', 1),
    ('lucknowi', 'Lucknowi', 'Subtle and aromatic', 'Yakhni-led', 'Mild', 1),
    ('hyderabadi', 'Hyderabadi', 'Bold and masaledar', 'Masala-led', 'Medium to hot', 3),
    ('mumbai', 'Mumbai', 'Familiar and masaledar', 'Mumbai-style flavour', 'Medium', 2),
    ('other-biryani', 'More biryani', 'Dishes outside the four regional styles', None, None, 0),
    ('tawa', 'Tawa Pulao & Anda Rice', 'Tossed to order on the tawa', None, None, 0),
    ('kebabs', 'Kebabs & Starters', None, None, None, 0),
    ('rolls', 'Rolls', None, None, None, 0),
    ('combos', 'Combos', 'Biryani, kebabs, gulab jamun and a drink in one order', None, None, 0),
    ('drinks', 'Drinks', None, None, None, 0),
]

PHONE = '+91 96196 11561'
ADDRESS = 'Second Floor, A-Wing, Express Zone, Malad East, Mumbai 400097'
HOURS = '11 AM - 11 PM, every day'
SIZES = [('half', 'Half'), ('full', 'Full'), ('perKg', '1 kg')]

PORTIONS = [
    ('Half', '350-375 g served  |  450 ml container  |  1 drumstick'),
    ('Full', '550-575 g served  |  750 ml container  |  2 drumsticks'),
    ('1 kg (Lucknowi)', '10 drumsticks  |  about 2.8-3.0 kg cooked  |  serves about 5-7'),
    ('1 kg curry cut', '18-20 pieces  |  no neck, no rib cage'),
]
ADDONS = [
    ('Mirchi ka Salan', 'Available with every biryani style, so you set your own heat.'),
    ('Extra masala', 'Hyderabadi only, on selected 1 kg bulk orders, with extra onion and tomato masala.'),
    ('Bulk additions', 'Kaju, mushroom and bell pepper on selected bulk orders. Ask us while ordering.'),
]

rs = lambda v: f'₹{v:,}'


def section_of(it):
    if it['category'] in ('biryani', 'pulao'):
        return it.get('region') or 'other-biryani'
    if it['category'] == 'tawa':
        return 'tawa'
    return it['category']


def variant_labels(variants):
    """Label rows by whatever actually differs: the cut, or the regional style."""
    if len(variants) < 2:
        return None  # the cut is already part of the dish name
    cuts = [v.get('cut') or '' for v in variants]
    if len(set(cuts)) > 1:
        return [c or 'Standard' for c in cuts]
    regions = [v.get('region') or '' for v in variants]
    if len(set(regions)) > 1:
        titles = {k: t for k, t, *_ in SECTIONS}
        return [titles.get(r, 'Standard') for r in regions]
    return cuts if cuts[0] else None


# ---------------------------------------------------------------- styles
W, H = A4
M = 17 * mm
CONTENT_W = W - 2 * M

dish = ParagraphStyle('dish', fontName='Mukta-SB', fontSize=10.5, leading=13.5, textColor=INK)
note = ParagraphStyle('note', fontName='Mukta', fontSize=8.5, leading=11, textColor=MUTED)
cut_st = ParagraphStyle('cut', fontName='Mukta', fontSize=9.5, leading=12.5, textColor=HexColor('#4E5C54'))
price = ParagraphStyle('price', fontName='Mukta-B', fontSize=10.5, leading=13.5, textColor=GREEN, alignment=TA_RIGHT)
na = ParagraphStyle('na', fontName='Mukta', fontSize=10, leading=13.5, textColor=FAINT, alignment=TA_RIGHT)
colhead = ParagraphStyle('ch', fontName='Mukta-SB', fontSize=7.5, leading=9, textColor=BRASS_DK, alignment=TA_RIGHT)
sec_title = ParagraphStyle('st', fontName='Rozha', fontSize=23, leading=26, textColor=GREEN)
sec_sub = ParagraphStyle('ss', fontName='Mukta', fontSize=9, leading=12, textColor=MUTED)
band = ParagraphStyle('band', fontName='Mukta', fontSize=9, leading=12.5, textColor=INK)
band_t = ParagraphStyle('bt', fontName='Mukta-SB', fontSize=9.5, leading=12.5, textColor=GREEN)


def dot(it):
    """The green / red square used on Indian menus."""
    if 'veg' not in it:
        return ''
    c = '#2E8B3E' if it['veg'] else '#B3261E'
    return f'<font name="ZapfDingbats" size="6.5" color="{c}">n</font>&nbsp;&nbsp;'


# ---------------------------------------------------------------- page furniture
def cover(c, doc):
    c.saveState()
    c.setFillColor(GREEN); c.rect(0, 0, W, H, stroke=0, fill=1)

    # jaali lattice, drawn faintly across the page
    c.setStrokeColor(BRASS); c.setLineWidth(0.4); c.setStrokeAlpha(0.13)
    step = 20 * mm
    y = 0
    while y < H + step:
        x = 0
        while x < W + step:
            c.lines([(x, y + step / 2, x + step / 2, y), (x + step / 2, y, x + step, y + step / 2),
                     (x + step, y + step / 2, x + step / 2, y + step), (x + step / 2, y + step, x, y + step / 2)])
            x += step
        y += step
    c.setStrokeAlpha(1)

    logo = os.path.join(ROOT, 'public', 'logo.png')
    c.drawImage(logo, W / 2 - 24 * mm, H - 86 * mm, 48 * mm, 48 * mm, mask='auto')

    c.setFillColor(IVORY); c.setFont('Rozha', 42)
    c.drawCentredString(W / 2, H - 104 * mm, "Daddu's Biryani")
    c.setFillColor(BRASS); c.setFont('Mukta-SB', 13)
    c.drawCentredString(W / 2, H - 113 * mm, 'Zayqo ki Kahani')

    c.setStrokeColor(BRASS); c.setLineWidth(0.8)
    c.line(W / 2 - 30 * mm, H - 122 * mm, W / 2 + 30 * mm, H - 122 * mm)

    c.setFillColor(IVORY); c.setFont('Rozha', 26)
    c.drawCentredString(W / 2, H - 140 * mm, 'Not one biryani.')
    c.drawCentredString(W / 2, H - 152 * mm, 'Many traditions.')

    c.setFillColor(BRASS); c.setFont('Mukta-SB', 12)
    c.drawCentredString(W / 2, H - 166 * mm, 'Kolkata   ·   Lucknowi   ·   Hyderabadi   ·   Mumbai')

    c.setFillColor(IVORY); c.setFont('Mukta', 11)
    c.drawCentredString(W / 2, H - 182 * mm, 'Every palate is different. Discover which biryani is yours.')
    c.setFillColor(BRASS); c.setFont('Rozha', 18)
    c.drawCentredString(W / 2, H - 195 * mm, 'Taste before you choose.')

    c.setFillColor(IVORY); c.setFont('Mukta', 10)
    c.drawCentredString(W / 2, 32 * mm, ADDRESS)
    c.drawCentredString(W / 2, 26 * mm, f'{HOURS}   |   {PHONE}  (call or WhatsApp)')
    c.restoreState()


def inner(c, doc):
    c.saveState()
    c.setFillColor(IVORY); c.rect(0, 0, W, H, stroke=0, fill=1)

    c.setFillColor(GREEN); c.rect(0, H - 15 * mm, W, 15 * mm, stroke=0, fill=1)
    c.setFillColor(IVORY); c.setFont('Rozha', 13)
    c.drawString(M, H - 10 * mm, "Daddu's Biryani")
    c.setFillColor(BRASS); c.setFont('Mukta-SB', 9.5)
    c.drawRightString(W - M, H - 10 * mm, f'{PHONE}  ·  {ADDRESS.split(",")[-2].strip()}')

    c.setStrokeColor(RULE); c.setLineWidth(0.6)
    c.line(M, 14 * mm, W - M, 14 * mm)
    c.setFillColor(MUTED); c.setFont('Mukta', 8)
    c.drawString(M, 9.5 * mm, 'Prices in INR. Taxes as applicable. NA means that size is not offered.')
    c.drawRightString(W - M, 9.5 * mm, f'Page {doc.page - 1}')

    c.setFillColor(VEG); c.rect(M, 4.2 * mm, 2.4 * mm, 2.4 * mm, stroke=0, fill=1)
    c.setFillColor(MUTED); c.setFont('Mukta', 8)
    c.drawString(M + 3.8 * mm, 4.4 * mm, 'Veg')
    c.setFillColor(NONVEG); c.rect(M + 13 * mm, 4.2 * mm, 2.4 * mm, 2.4 * mm, stroke=0, fill=1)
    c.setFillColor(MUTED); c.drawString(M + 16.8 * mm, 4.4 * mm, 'Non-veg')
    c.restoreState()


def chillies(level):
    if not level:
        return ''
    on = '<font color="#B3261E">●</font>' * level
    off = f'<font color="{FAINT.hexval()[2:]}">○</font>' * (3 - level)
    return on + off


def section_head(title, taste, character, spice, level):
    bits = []
    if taste:
        bits.append(taste)
    if character:
        bits.append(character)
    if spice:
        bits.append(f'Spice level: {spice}')
    sub = '   ·   '.join(bits)
    rows = [[Paragraph(title, sec_title)]]
    if sub:
        rows.append([Paragraph(sub, sec_sub)])
    t = Table(rows, colWidths=[CONTENT_W])
    t.setStyle(TableStyle([
        ('LINEABOVE', (0, 0), (-1, 0), 1.1, GREEN),
        ('TOPPADDING', (0, 0), (-1, 0), 7), ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ('LEFTPADDING', (0, 0), (-1, -1), 0), ('RIGHTPADDING', (0, 0), (-1, -1), 0),
    ]))
    return t


# ---------------------------------------------------------------- build the story
story = [NextPageTemplate('inner'), PageBreak()]

for key, title, taste, character, spice, level in SECTIONS:
    rows_in = [i for i in items if section_of(i) == key]
    if not rows_in:
        continue

    groups = []
    for it in rows_in:
        g = next((x for x in groups if x[0] == it['name']), None)
        if g:
            g[1].append(it)
        else:
            groups.append((it['name'], [it]))

    cols = [(k, lbl) for k, lbl in SIZES if any(v['prices'].get(k) for _, vs in groups for v in vs)]
    simple = not cols  # single-price sections print in two columns

    head = KeepTogether([CondPageBreak(48 * mm), section_head(title, taste, character, spice, level), Spacer(1, 4 * mm)])

    if simple:
        # two columns of "name .......... price"
        cells = []
        for name, variants in groups:
            v = variants[0]
            txt = dot(v) + name
            extra = ''
            if v.get('includes'):
                extra = ', '.join(v['includes'])
                if v.get('serves'):
                    extra += f" · serves {v['serves']}"
            elif v.get('note'):
                extra = v['note']
            inner_rows = [[Paragraph(txt, dish), Paragraph(rs(v['prices'].get('price') or v['prices'].get('half')), price)]]
            if extra:
                inner_rows.append([Paragraph(extra, note), ''])
            t = Table(inner_rows, colWidths=[(CONTENT_W / 2 - 7 * mm) - 17 * mm, 17 * mm])
            t.setStyle(TableStyle([
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('SPAN', (0, 1), (1, 1)) if extra else ('NOSPAN', (0, 0), (0, 0)),
                ('LEFTPADDING', (0, 0), (-1, -1), 0), ('RIGHTPADDING', (0, 0), (-1, -1), 0),
                ('TOPPADDING', (0, 0), (-1, -1), 0), ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ]))
            cells.append(t)

        half = (len(cells) + 1) // 2
        left, right = cells[:half], cells[half:]
        right += [''] * (len(left) - len(right))
        grid = Table([[l, r] for l, r in zip(left, right)],
                     colWidths=[CONTENT_W / 2 - 7 * mm, CONTENT_W / 2 - 7 * mm])
        grid.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (0, -1), 0), ('RIGHTPADDING', (0, 0), (0, -1), 14),
            ('LEFTPADDING', (1, 0), (1, -1), 14), ('RIGHTPADDING', (1, 0), (1, -1), 0),
            ('TOPPADDING', (0, 0), (-1, -1), 2), ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story += [head, grid, Spacer(1, 7 * mm)]
        continue

    # sized sections: name on the left, aligned price columns on the right
    pw = 20 * mm
    widths = [CONTENT_W - pw * len(cols)] + [pw] * len(cols)
    data = [[''] + [Paragraph(lbl, colhead) for _, lbl in cols]]
    styles = [
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LINEBELOW', (0, 0), (-1, 0), 0.5, RULE),
        ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 0), ('RIGHTPADDING', (0, 0), (-1, -1), 0),
    ]

    for name, variants in groups:
        lead = variants[0]
        labels = variant_labels(variants)
        r = len(data)

        price_cells = []
        if not labels:
            if lead['prices'].get('price'):
                price_cells = [Paragraph(rs(lead['prices']['price']), price)] + [''] * (len(cols) - 1)
                styles.append(('SPAN', (1, r), (len(cols), r)))
            else:
                price_cells = [Paragraph(rs(lead['prices'][k]), price) if lead['prices'].get(k)
                               else Paragraph('NA', na) for k, _ in cols]
        else:
            price_cells = [''] * len(cols)

        data.append([Paragraph(dot(lead) + name, dish)] + price_cells)
        styles.append(('LINEABOVE', (0, r), (-1, r), 0.4, HexColor('#EFE6D2')))

        extra = lead.get('note') or ''
        if lead.get('includes'):
            extra = ', '.join(lead['includes'])
        if extra:
            data.append([Paragraph(extra, note)] + [''] * len(cols))
            styles.append(('TOPPADDING', (0, len(data) - 1), (-1, len(data) - 1), 0))

        if labels:
            for v, lbl in zip(variants, labels):
                cells = [Paragraph(rs(v['prices'][k]), price) if v['prices'].get(k)
                         else Paragraph('NA', na) for k, _ in cols]
                data.append([Paragraph(f'&nbsp;&nbsp;&nbsp;&nbsp;{lbl}', cut_st)] + cells)

    t = Table(data, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle(styles))
    story += [head, t, Spacer(1, 7 * mm)]

# ---- closing bands: portions and add-ons -------------------------------------
def band_table(title, rows):
    inner_rows = [[Paragraph(k, band_t), Paragraph(v, band)] for k, v in rows]
    t = Table(inner_rows, colWidths=[38 * mm, CONTENT_W - 38 * mm - 20])
    t.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('BACKGROUND', (0, 0), (-1, -1), IVORY_DIM),
        ('LEFTPADDING', (0, 0), (0, -1), 10), ('RIGHTPADDING', (-1, 0), (-1, -1), 10),
        ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    return KeepTogether([CondPageBreak(42 * mm), section_head(title, None, None, None, 0), Spacer(1, 3 * mm), t, Spacer(1, 7 * mm)])


story.append(band_table('What you actually get', PORTIONS))
story.append(band_table('Add-ons and customisation', ADDONS))

closing = Table([[Paragraph(
    "<font name='Rozha' size='15' color='#0E3B2F'>Don't just order biryani. Discover your biryani.</font><br/>"
    f"<font name='Mukta' size='10' color='#6B7A71'>Call or WhatsApp {PHONE} · {ADDRESS}</font>",
    ParagraphStyle('c', alignment=TA_CENTER, leading=20))]], colWidths=[CONTENT_W])
closing.setStyle(TableStyle([
    ('LINEABOVE', (0, 0), (-1, 0), 1.1, BRASS),
    ('TOPPADDING', (0, 0), (-1, -1), 10), ('LEFTPADDING', (0, 0), (-1, -1), 0), ('RIGHTPADDING', (0, 0), (-1, -1), 0),
]))
story.append(closing)

pdf_path = os.path.join(OUT, 'daddus-biryani-menu.pdf')
doc = BaseDocTemplate(pdf_path, pagesize=A4, title="Daddu's Biryani Menu",
                      author="Daddu's Biryani", subject='Menu and prices')
doc.addPageTemplates([
    PageTemplate('cover', [Frame(M, 18 * mm, CONTENT_W, H - 40 * mm, id='cv')], onPage=cover),
    PageTemplate('inner', [Frame(M, 18 * mm, CONTENT_W, H - 15 * mm - 24 * mm, id='in')], onPage=inner),
])
doc.build(story)

# ---------------------------------------------------------------- Excel
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Menu'
thin = Side(style='thin', color='E0D4BA')

ws['A1'] = "Daddu's Biryani - Menu"
ws['A1'].font = Font(name='Arial', size=18, bold=True, color='0E3B2F')
ws['A2'] = 'Not one biryani. Many traditions.  |  Taste before you choose.'
ws['A2'].font = Font(name='Arial', size=11, color='9A7A2E')
ws['A3'] = f'{PHONE}   |   {ADDRESS}   |   {HOURS}'
ws['A3'].font = Font(name='Arial', size=10, color='6B7A71')
ws['A4'] = 'Prices in INR. Taxes as applicable. Blank means that size is not offered.'
ws['A4'].font = Font(name='Arial', size=9, italic=True, color='6B7A71')

hdr = ['Section', 'Item', 'Variant', 'Veg / Non-veg', 'Half', 'Full', 'Per kg', 'Price', 'Details']
r = 6
for c, h in enumerate(hdr, 1):
    cell = ws.cell(r, c, h)
    cell.font = Font(name='Arial', bold=True, color='FBF5E9')
    cell.fill = PatternFill('solid', fgColor='0E3B2F')
    cell.alignment = Alignment(vertical='center')

titles = {k: t for k, t, *_ in SECTIONS}
for key, *_ in SECTIONS:
    rows_in = [i for i in items if section_of(i) == key]
    groups = []
    for it in rows_in:
        g = next((x for x in groups if x[0] == it['name']), None)
        if g:
            g[1].append(it)
        else:
            groups.append((it['name'], [it]))

    for name, variants in groups:
        labels = variant_labels(variants) or [''] * len(variants)
        for v, lbl in zip(variants, labels):
            r += 1
            p = v['prices']
            details = v.get('note') or ''
            if v.get('includes'):
                details = ', '.join(v['includes'])
                if v.get('serves'):
                    details += f" (serves {v['serves']})"
            vals = [titles[key], name, lbl,
                    '' if 'veg' not in v else ('Veg' if v['veg'] else 'Non-veg'),
                    p.get('half'), p.get('full'), p.get('perKg'), p.get('price'), details]
            for c, val in enumerate(vals, 1):
                cell = ws.cell(r, c, val)
                cell.font = Font(name='Arial', size=10, color='1F2A24')
                cell.border = Border(bottom=thin)
                if 5 <= c <= 8:
                    cell.number_format = '#,##0'
                    cell.alignment = Alignment(horizontal='right')
                if c == 9:
                    cell.alignment = Alignment(wrap_text=True, vertical='top')

last = r
r += 2
ws.cell(r, 1, 'What you actually get').font = Font(name='Arial', bold=True, size=11, color='0E3B2F')
for k, v in PORTIONS:
    r += 1
    ws.cell(r, 1, k).font = Font(name='Arial', size=10, bold=True, color='0E3B2F')
    ws.cell(r, 2, v).font = Font(name='Arial', size=10, color='1F2A24')
r += 2
ws.cell(r, 1, 'Add-ons and customisation').font = Font(name='Arial', bold=True, size=11, color='0E3B2F')
for k, v in ADDONS:
    r += 1
    ws.cell(r, 1, k).font = Font(name='Arial', size=10, bold=True, color='0E3B2F')
    ws.cell(r, 2, v).font = Font(name='Arial', size=10, color='1F2A24')

for col, w in zip('ABCDEFGHI', [22, 36, 14, 13, 9, 9, 9, 9, 58]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A7'
ws.auto_filter.ref = f'A6:I{last}'

xlsx_path = os.path.join(OUT, 'daddus-biryani-menu.xlsx')
wb.save(xlsx_path)
print('PDF ', pdf_path, os.path.getsize(pdf_path) // 1024, 'KB')
print('XLSX', xlsx_path, os.path.getsize(xlsx_path) // 1024, 'KB')
